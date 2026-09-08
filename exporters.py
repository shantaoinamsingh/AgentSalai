#!/usr/bin/env python3
"""Turn assistant output into downloadable files.

The browser parses a rendered Markdown table (or a chart's dataset) into
``{"headers": [...], "rows": [[...]]}`` and posts it here. Keeping the
conversion server-side means one implementation for CSV and XLSX, and real
number typing in the spreadsheet rather than text-formatted digits.

PDF is deliberately *not* here: it is produced in the browser via the print
pipeline, which reproduces the rendered Markdown, tables and canvas charts
exactly as the user sees them. A server-side text-only PDF would look nothing
like the answer on screen, and the alternative engines are heavy dependencies
(WeasyPrint needs GTK on Windows).
"""
import csv
import io
import re
from datetime import datetime, timezone
from typing import Any, Dict, List, Optional, Tuple

MAX_ROWS = 50_000
MAX_COLS = 256


class ExportError(ValueError):
    """The submitted table could not be exported."""


def safe_filename(stem: str, extension: str) -> str:
    """A download name that is safe on Windows, macOS and Linux."""
    stem = re.sub(r"[^A-Za-z0-9 ._-]+", "_", (stem or "").strip())
    stem = re.sub(r"\s+", "_", stem).strip("._") or "salai-export"
    stem = stem[:60]
    stamp = datetime.now(timezone.utc).strftime("%Y%m%d-%H%M%S")
    return f"{stem}-{stamp}.{extension.lstrip('.')}"


def normalise_table(payload: Dict[str, Any]) -> Tuple[List[str], List[List[Any]]]:
    """Validate an incoming table and pad it into a clean rectangle."""
    if not isinstance(payload, dict):
        raise ExportError("Expected a JSON object")

    headers = payload.get("headers") or []
    rows = payload.get("rows") or []

    if not isinstance(headers, list) or not isinstance(rows, list):
        raise ExportError("'headers' and 'rows' must both be arrays")
    if not headers and not rows:
        raise ExportError("There is no table data to export")
    if len(rows) > MAX_ROWS:
        raise ExportError(f"Too many rows to export (limit {MAX_ROWS:,})")

    headers = [str(h) for h in headers[:MAX_COLS]]

    width = max([len(headers)] + [len(r) for r in rows if isinstance(r, list)] or [0])
    width = min(width, MAX_COLS)
    if not headers:
        headers = [f"Column {i + 1}" for i in range(width)]
    headers += [""] * (width - len(headers))

    clean: List[List[Any]] = []
    for row in rows:
        if not isinstance(row, list):
            raise ExportError("Every row must be an array")
        cells = list(row[:width]) + [""] * (width - len(row))
        clean.append([_coerce(c) for c in cells])
    return headers, clean


def _coerce(cell: Any) -> Any:
    """Store numbers as numbers so spreadsheets can chart and total them."""
    if cell is None:
        return ""
    if isinstance(cell, (int, float, bool)):
        return cell
    text = str(cell).strip()
    if not text:
        return ""

    # Tolerate thousands separators, currency symbols, trailing % and
    # parenthesised negatives, all of which models emit in tables.
    candidate = text.replace(",", "").replace("£", "").replace("$", "").replace("€", "")
    negative = candidate.startswith("(") and candidate.endswith(")")
    if negative:
        candidate = candidate[1:-1]
    percent = candidate.endswith("%")
    if percent:
        candidate = candidate[:-1]

    try:
        value = float(candidate)
    except ValueError:
        return text

    if negative:
        value = -value
    if percent:
        # Keep the sign in the text form; a bare number would lose the unit.
        return text
    return int(value) if value.is_integer() and abs(value) < 1e15 else value


def to_csv(headers: List[str], rows: List[List[Any]]) -> bytes:
    buffer = io.StringIO(newline="")
    writer = csv.writer(buffer, quoting=csv.QUOTE_MINIMAL, lineterminator="\r\n")
    writer.writerow(headers)
    writer.writerows(rows)
    # BOM so Excel opens UTF-8 CSVs with the right encoding on Windows.
    return b"\xef\xbb\xbf" + buffer.getvalue().encode("utf-8")


def to_xlsx(headers: List[str], rows: List[List[Any]], title: str = "Sheet1") -> bytes:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        raise ExportError(
            "Excel export needs openpyxl. Run: pip install openpyxl (or download CSV)"
        )

    workbook = Workbook()
    sheet = workbook.active
    # Excel sheet names cap at 31 chars and reject : \ / ? * [ ]
    sheet.title = (re.sub(r"[:\\/?*\[\]]", "-", title or "Sheet1") or "Sheet1")[:31]

    sheet.append(headers)
    header_fill = PatternFill("solid", fgColor="667EEA")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")

    for row in rows:
        sheet.append(row)

    # Width from the longest value in each column, within sane bounds.
    for index in range(1, len(headers) + 1):
        longest = len(str(headers[index - 1]))
        for row in rows[:200]:
            if index - 1 < len(row):
                longest = max(longest, len(str(row[index - 1])))
        sheet.column_dimensions[get_column_letter(index)].width = min(max(longest + 2, 9), 55)

    sheet.freeze_panes = "A2"
    if rows:
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"

    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def build(fmt: str, payload: Dict[str, Any]) -> Tuple[bytes, str, str]:
    """Render a table export. Returns ``(bytes, mimetype, filename)``."""
    headers, rows = normalise_table(payload)
    title: Optional[str] = payload.get("title") or "salai-export"

    if fmt == "csv":
        return (
            to_csv(headers, rows),
            "text/csv; charset=utf-8",
            safe_filename(title, "csv"),
        )
    if fmt in ("xlsx", "excel"):
        return (
            to_xlsx(headers, rows, title),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            safe_filename(title, "xlsx"),
        )
    raise ExportError(f"Unsupported export format: {fmt!r}. Use csv or xlsx.")
